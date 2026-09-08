"""The 56 local cases prove conversion, not any platform's cloud read/write compatibility."""

import json
import zipfile
from pathlib import Path
from urllib.parse import unquote

import pytest
from lxml import etree, html

from note_bridge.errors import Cancelled
from note_bridge.exporter import Exporter, package_export
from note_bridge.models import Attachment, Block, PlatformId, Span
from note_bridge.paths import safe_filename
from note_bridge.richtext import blocks_html, markdown_span, parse_html


@pytest.mark.parametrize("blank_lines", [1, 2, 3])
def test_html_roundtrip_preserves_intentional_blank_paragraphs(blank_lines):
    original = [Block(spans=[Span(text="before")]), *[Block() for _ in range(blank_lines)],
                Block(spans=[Span(text="after", bold=True)])]
    restored, _ = parse_html(blocks_html(original, {}))
    assert restored == original, "空白段落是正文的一部分，迁移和离线读取不能合并它们"


def test_inline_break_and_container_spacing_are_not_extra_blank_paragraphs():
    restored, _ = parse_html('<div>\n<p>before<br>after</p>\n</div>')
    assert restored == [Block(spans=[Span(text="before\nafter")])]


@pytest.mark.parametrize("platform", list(PlatformId))
@pytest.mark.parametrize("format", ["txt", "md", "html", "docx"])
@pytest.mark.parametrize("multi", [False, True])
def test_56_offline_exports_preserve_documents_and_portable_resources(
    tmp_path, corpus, platform, format, multi
):
    resources, originals = corpus
    notes = [note.model_copy(update={"platform": platform}, deep=True) for note in originals]
    result = Exporter(resources).export(notes, tmp_path, format, multi)
    files = list(result.path.glob(f"*.{format}"))
    assert len(files) == (4 if multi else 1), "重名和空标题不能覆盖或漏掉笔记"
    assert len(list((result.path / "resources").iterdir())) == 2
    report = json.loads((result.path / "导出说明.json").read_text(encoding="utf8"))
    assert report["note_count"] == 4
    if format == "docx":
        with zipfile.ZipFile(files[0]) as archive:
            xml = etree.fromstring(archive.read("word/document.xml"))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            text = "".join(xml.xpath("//w:t/text()", namespaces=ns))
            assert "中文 English 😀" in text and "2026-08-15" in text
            assert xml.xpath("//w:b", namespaces=ns) and xml.xpath("//w:tbl", namespaces=ns)
            assert any(name.startswith("word/media/") for name in archive.namelist())
            if not multi:
                anchors = set(xml.xpath("//w:hyperlink/@w:anchor", namespaces=ns))
                bookmarks = set(xml.xpath("//w:bookmarkStart/@w:name", namespaces=ns))
                assert len(anchors) == 4 and anchors <= bookmarks, "合并 DOCX 目录的每一项必须可跳转"
    elif format == "html":
        document = html.fromstring(files[0].read_text(encoding="utf8"))
        if multi:
            content = document
        else:
            records = json.loads(document.get_element_by_id("note-data").text)
            assert len(records) == 4
            content = html.fromstring(records[0]["html"])
        assert "中文 English 😀" in content.text_content()
        assert content.xpath("//strong") and content.xpath("//table")
        assert content.xpath("//ol[@start='2']"), "有序列表的第二项不能重新编号为 1"
        for src in content.xpath("//img/@src"):
            assert (result.path / unquote(src)).is_file()
    else:
        text = "\n".join(path.read_text(encoding="utf-8-sig") for path in files)
        assert "中文 English 😀" in text and "2026-08-15" in text
        assert "长文本" * 1500 in text and "无标题" in text
        assert "2. second" in text
    assert not list(tmp_path.glob(".note-bridge-*")), "完成后不应残留中间目录"


@pytest.mark.parametrize("multi", [False, True])
def test_docx_links_preserve_combined_styles_and_inline_structure(tmp_path, corpus, multi):
    from docx.text.run import Run

    resources, notes = corpus
    note = notes[0].model_copy(deep=True)
    note.attachments = []
    url = "https://example.com/notes?q=one&lang=zh"
    styled = Span(text="  链接😀\n第二行\t末尾  ", link=url, bold=True, italic=True,
                  underline=True, strike=True, highlight=True, code=True)
    plain = Span(text="普通链接", link=url)
    note.blocks = [Block(spans=[styled]), Block(spans=[plain]), Block(kind="code", spans=[plain])]
    result = Exporter(resources).export([note], tmp_path, "docx", multi)
    from docx import Document

    doc = Document(next(result.path.glob("*.docx")))
    links = doc.element.xpath('//w:hyperlink[@r:id]')
    assert len(links) == 3
    for link in links:
        assert doc.part.rels[link.get('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id')].target_ref == url
    styled_run, plain_run, code_run = [Run(link[0], doc) for link in links]
    assert styled_run.text == styled.text, "命名链接内的空格、换行和制表符也是正文，不能丢失"
    assert styled_run.bold and styled_run.italic and styled_run.underline and styled_run.font.strike
    assert styled_run.font.highlight_color is not None and styled_run.font.name == "Consolas"
    assert not any([plain_run.bold, plain_run.italic, plain_run.underline, plain_run.font.strike,
                    plain_run.font.highlight_color]), "相邻普通链接不能继承前一个片段的格式"
    assert code_run.font.name == "Consolas", "代码块中的链接仍应保留代码字体"
    assert len(styled_run.element.xpath('./w:br')) == len(styled_run.element.xpath('./w:tab')) == 1


@pytest.mark.parametrize("multi", [False, True])
def test_docx_numbering_preserves_nested_sequences_and_restarts_between_notes(tmp_path, corpus, multi):
    from docx import Document
    from docx.oxml.ns import qn

    resources, notes = corpus
    first, second = [n.model_copy(deep=True) for n in notes[:2]]
    first.attachments = []
    first.blocks = [Block(kind="list", ordered=True, level=level, spans=[Span(text=text)])
                    for level, text in [(1, "top-a"), (2, "sub-a"), (2, "sub-b"), (1, "top-b")]]
    first.blocks += [Block(spans=[Span(text="separate")]),
                     Block(kind="list", ordered=True, spans=[Span(text="top-c")])]
    second.blocks = [Block(kind="list", ordered=True, spans=[Span(text="next-note")])]
    result = Exporter(resources).export([first, second], tmp_path, "docx", multi)
    observed = {}
    for path in sorted(result.path.glob('*.docx')):
        doc = Document(path)
        numbering = doc.part.numbering_part.element
        counters = {}
        for paragraph in doc.paragraphs:
            if paragraph.text not in {'top-a', 'sub-a', 'sub-b', 'top-b', 'top-c', 'next-note'}:
                continue
            properties = paragraph._p.pPr.numPr
            if properties is None:
                properties = paragraph.style.element.pPr.numPr
            num_id = properties.numId.val
            definition = numbering.num_having_numId(num_id)
            starts = definition.xpath('./w:lvlOverride[@w:ilvl="0"]/w:startOverride')
            abstract_id = definition.abstractNumId.val
            if not starts:
                starts = numbering.xpath(f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="0"]/w:start')
            counters[num_id] = counters.get(num_id, int(starts[0].get(qn('w:val'))) - 1) + 1
            observed[paragraph.text] = counters[num_id]
    assert observed == {'top-a': 1, 'sub-a': 1, 'sub-b': 2, 'top-b': 2, 'top-c': 1, 'next-note': 1}, (
        "嵌套编号独立递增，返回上层继续；普通段落和下一条笔记开始新列表，不能沿用全篇编号"
    )


def test_cancel_does_not_remove_previous_exports_or_user_files(tmp_path, corpus):
    resources, notes = corpus
    sentinel = tmp_path / "用户已有文件.txt"
    sentinel.write_text("keep", encoding="utf8")
    calls = 0

    def cancel():
        nonlocal calls
        calls += 1
        if calls == 3:
            raise Cancelled()

    with pytest.raises(Cancelled):
        Exporter(resources).export(notes, tmp_path, "html", True, check_cancel=cancel)
    assert sentinel.read_text(encoding="utf8") == "keep"
    assert not list(tmp_path.glob(".note-bridge-*"))


def test_missing_corrupt_and_unsupported_images_are_reported(tmp_path, corpus):
    resources, notes = corpus
    notes[0].attachments[0].sha256 = "0" * 64
    notes[0].attachments.append(Attachment(id="missing", name="missing.pdf"))
    result = Exporter(resources).export(notes, tmp_path, "html")
    assert {issue.code for issue in result.issues} >= {"attachment_corrupt", "attachment_missing"}
    notes[0].attachments[0].sha256 = None
    (resources / "示例.png").write_bytes(b"this is not an image")
    result = Exporter(resources).export(notes, tmp_path, "docx")
    assert "image_not_embedded" in {issue.code for issue in result.issues}
    assert any(path.read_bytes() == b"this is not an image" for path in (result.path / "resources").iterdir())


def test_bundle_has_only_relative_entries_and_survives_relocation(tmp_path, corpus):
    resources, notes = corpus
    result = Exporter(resources).export(notes, tmp_path, "html")
    zip_path = package_export(result.path, tmp_path)
    with zipfile.ZipFile(zip_path) as archive:
        assert all(
            not name.startswith(("/", "\\")) and ".." not in Path(name).parts for name in archive.namelist()
        )
        assert len(archive.namelist()) == 4
        relocated = tmp_path / "换一台电脑"
        archive.extractall(relocated)
    document = html.fromstring(next(relocated.glob("*.html")).read_text(encoding="utf8"))
    record = json.loads(document.get_element_by_id("note-data").text)[0]
    content = html.fromstring(record["html"])
    assert (relocated / unquote(content.xpath("//img/@src")[0])).is_file()


def test_executable_notes_never_become_html_or_markdown_code():
    blocks, warnings = parse_html(
        '<p onclick="alert(1)">你好<b>加粗</b><a href="javascript:alert(1)">危险链接</a></p><script>alert(1)</script><iframe src="https://example.com"></iframe>'
    )
    rendered = blocks_html(blocks, {})
    assert warnings
    assert all(value not in rendered for value in ("<script", "<iframe", "onclick", "javascript:"))
    assert "你好" in rendered and "<strong>加粗</strong>" in rendered
    span = Span(text="<img src=x onerror=alert(1)>", underline=True)
    assert "<img" not in markdown_span(span)


@pytest.mark.parametrize("name", ["CON", "aux.txt", "../x", "a:b?*<>|", " ", "NUL."])
def test_windows_names_cannot_escape_or_target_devices(name):
    safe = safe_filename(name)
    assert safe and not any(c in safe for c in '<>:"/\\|?*')
    assert safe.upper().split(".")[0] not in ("CON", "AUX", "NUL")


def test_attachment_path_cannot_read_outside_resources(tmp_path, corpus):
    resources, notes = corpus
    (tmp_path / "private.txt").write_text("do not export", encoding="utf8")
    notes[0].attachments = [Attachment(id="file", name="private.txt", local_path="../private.txt")]
    result = Exporter(resources).export(notes, tmp_path, "txt")
    assert "attachment_missing" in {i.code for i in result.issues}
    assert not (result.path / "resources").exists()


def test_transient_windows_directory_lock_retries_only_the_atomic_commit(monkeypatch, tmp_path, corpus):
    from note_bridge import exporter

    resources, notes = corpus
    original = exporter.os.replace
    calls = []

    def transient(source, target):
        calls.append((source, target))
        if len(calls) <= 2:
            error = PermissionError("synthetic directory lock")
            error.winerror = 5
            raise error
        return original(source, target)

    monkeypatch.setattr(exporter.os, "replace", transient)
    result = Exporter(resources).export(notes, tmp_path, "md")
    assert len(calls) == 3 and len(set(calls)) == 1
    assert result.note_count == len(notes) and len(list(result.path.glob("*.md"))) == 1
