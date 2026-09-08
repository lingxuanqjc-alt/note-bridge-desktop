"""Atomic, portable exports. A missing resource is always a reported issue."""

from __future__ import annotations

import hashlib
import html
import json
import os
import shutil
import time
import uuid
import zipfile
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Callable, Literal
from urllib.parse import quote

from docx import Document
from docx.enum.text import WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from docx.text.run import Run

from .errors import BridgeError
from .models import PLATFORM_NAMES, ItemIssue, NoteDocument
from .paths import confined, safe_filename
from .reader import reader_html
from .richtext import block_html, blocks_html, blocks_markdown, numbered_blocks, safe_link

Format = Literal["txt", "md", "html", "docx"]


@dataclass
class ExportResult:
    path: Path
    note_count: int
    issues: list[ItemIssue]


def stamp(note: NoteDocument) -> str:
    created = note.created_at.isoformat(sep=" ", timespec="seconds") if note.created_at else "未知"
    updated = note.updated_at.isoformat(sep=" ", timespec="seconds") if note.updated_at else "未知"
    return f"创建时间：{created}\n修改时间：{updated}"


def unused_path(parent: Path, name: str) -> Path:
    path = parent / name
    if not path.exists():
        return path
    for index in range(2, 100_000):
        path = parent / f"{Path(name).stem} ({index}){Path(name).suffix}"
        if not path.exists():
            return path
    raise BridgeError("name_conflict", "输出目录中存在过多重名文件，请更换目录。")


def commit_directory(stage: Path, target: Path, check_cancel: Callable[[], None]):
    # Windows may briefly hold a newly written directory open (for example during indexing).
    # Retry only this atomic rename, never a cloud write or a partially repeated export.
    for attempt in range(4):
        check_cancel()
        try:
            os.replace(stage, target)
            return
        except PermissionError as error:
            if getattr(error, "winerror", None) not in (5, 32) or attempt == 3 or target.exists():
                raise
            time.sleep(0.05 * 2**attempt)


class Exporter:
    def __init__(self, resources: Path):
        self.resources = resources

    def export(
        self,
        notes: list[NoteDocument],
        destination: Path,
        format: Format,
        multi: bool = False,
        progress: Callable[[int, int], None] = lambda *_: None,
        check_cancel: Callable[[], None] = lambda: None,
    ) -> ExportResult:
        if format not in ("txt", "md", "html", "docx"):
            raise BridgeError("invalid_format", "请选择 TXT、Markdown、HTML 或 DOCX。")
        if not notes:
            raise BridgeError("empty_notes", "没有可以导出的笔记，请先获取笔记。")
        destination = destination.resolve()
        if not destination.is_dir():
            raise BridgeError("invalid_directory", "输出目录不存在，请重新选择。")
        label = PLATFORM_NAMES[notes[0].platform]
        stage = confined(destination, f".note-bridge-{uuid.uuid4().hex}")
        stage.mkdir()
        issues: list[ItemIssue] = []
        try:
            records, mappings = [], []
            for index, note in enumerate(notes):
                check_cancel()
                mapping = self._copy_assets(note, stage, issues)
                known_assets = {asset.id for asset in note.attachments}
                for block in note.blocks:
                    if block.kind == "attachment" and block.attachment_id not in known_assets:
                        issues.append(
                            ItemIssue(
                                note_id=note.source_id,
                                code="attachment_missing",
                                message="正文引用的附件缺少元数据，已保留缺失提示。",
                            )
                        )
                mappings.append(mapping)
                for warning in note.warnings:
                    issues.append(ItemIssue(note_id=note.source_id, code="source_warning", message=warning))
                if format == "txt" and any(
                    block.kind not in ("paragraph", "attachment")
                    or any(
                        span.bold
                        or span.italic
                        or span.underline
                        or span.strike
                        or span.highlight
                        or span.link
                        for span in block.spans
                    )
                    for block in note.blocks
                ):
                    issues.append(
                        ItemIssue(
                            note_id=note.source_id,
                            code="format_downgrade",
                            message="TXT 不支持富文本样式，已保留正文、列表标记和附件路径。",
                        )
                    )
                if format == "docx" and any(
                    asset.kind in ("audio", "video", "file") for asset in note.attachments
                ):
                    issues.append(
                        ItemIssue(
                            note_id=note.source_id,
                            code="format_downgrade",
                            message="音视频和其他文件作为相对链接保留在附件目录。",
                        )
                    )
                rendered = self._note_html(note, mapping)
                records.append(
                    {
                        "title": note.display_title,
                        "text": note.plain_text,
                        "summary": note.plain_text.replace("\n", " ")[:100],
                        "created": note.created_at.isoformat() if note.created_at else "",
                        "html": rendered,
                    }
                )
                if multi:
                    filename = f"{index + 1:04d}、{safe_filename(note.display_title)}.{format}"
                    self._write(
                        stage / filename, [note], [mapping], [records[-1]], format, True, label, issues
                    )
                progress(index + 1, len(notes))
            check_cancel()
            if not multi:
                self._write(
                    stage / f"{label}.{format}", notes, mappings, records, format, False, label, issues
                )
            # An export report travels with the files and contains no authentication material.
            report = {
                "application": "笔记互迁",
                "note_count": len(notes),
                "format": format,
                "mode": "multi" if multi else "single",
                "issues": [issue.model_dump() for issue in issues],
            }
            (stage / "导出说明.json").write_text(
                json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8"
            )
            check_cancel()
            final = unused_path(
                destination, f"{label}-{datetime.now():%Y%m%d-%H%M%S}-{format}{'-多文件' if multi else ''}"
            )
            commit_directory(stage, final, check_cancel)
            return ExportResult(final, len(notes), issues)
        except BaseException:
            # Only remove the exact newly-created staging directory within the selected export directory.
            resolved = stage.resolve()
            if (
                resolved.parent == destination
                and resolved.name.startswith(".note-bridge-")
                and resolved.is_dir()
            ):
                shutil.rmtree(resolved)
            raise

    def _copy_assets(
        self, note: NoteDocument, stage: Path, issues: list[ItemIssue]
    ) -> dict[str, tuple[str, str, str]]:
        result = {}
        for asset in note.attachments:
            if not asset.local_path:
                issues.append(
                    ItemIssue(
                        note_id=note.source_id, code="attachment_missing", message=f"附件未下载：{asset.name}"
                    )
                )
                continue
            try:
                source = confined(self.resources, asset.local_path)
                if not source.is_file():
                    raise OSError("missing")
                with source.open("rb") as stream:
                    digest = hashlib.file_digest(stream, "sha256").hexdigest()
                if asset.sha256 and asset.sha256 != digest:
                    issues.append(
                        ItemIssue(
                            note_id=note.source_id,
                            code="attachment_corrupt",
                            message=f"附件校验不通过：{asset.name}",
                        )
                    )
                    continue
                name = f"{digest[:16]}-{safe_filename(asset.name)}"
                target = stage / "resources" / name
                target.parent.mkdir(exist_ok=True)
                if not target.exists():
                    shutil.copyfile(source, target)
                result[asset.id] = ("resources/" + quote(name, safe="-._"), asset.name, asset.kind)
            except (OSError, BridgeError):
                issues.append(
                    ItemIssue(
                        note_id=note.source_id,
                        code="attachment_missing",
                        message=f"附件无法读取：{asset.name}",
                    )
                )
        return result

    def _note_html(self, note: NoteDocument, mapping: dict) -> str:
        content = f'<h2>{html.escape(note.display_title)}</h2><div class="meta">{html.escape(stamp(note)).replace(chr(10), "<br>")}</div>'
        content += blocks_html(note.blocks, mapping)
        used = {block.attachment_id for block in note.blocks if block.kind == "attachment"}
        from .models import Block

        for asset in note.attachments:
            if asset.id not in used:
                content += block_html(Block(kind="attachment", attachment_id=asset.id), mapping)
        for warning in note.warnings:
            content += f'<p class="warning">{html.escape(warning)}</p>'
        return content

    def _write(
        self,
        path: Path,
        notes: list[NoteDocument],
        mappings: list[dict],
        records: list[dict],
        format: str,
        single: bool,
        label: str,
        issues: list[ItemIssue],
    ):
        if format == "docx":
            self._docx(path, notes, mappings, not single, issues)
            return
        if format == "html":
            text = reader_html(label + " · 笔记互迁", records, single=single)
        elif format == "md":
            parts = []
            for note, mapping in zip(notes, mappings, strict=True):
                body = blocks_markdown(note.blocks, mapping)
                unplaced = [
                    asset
                    for asset in note.attachments
                    if asset.id not in {block.attachment_id for block in note.blocks}
                ]
                for asset in unplaced:
                    if asset.id in mapping:
                        body += f"\n\n[附件：{asset.name.replace(']', '')}]({mapping[asset.id][0]})"
                    else:
                        body += f"\n\n[附件未获取：{asset.name}]"
                parts.append(f"# {html.escape(note.display_title)}\n\n{stamp(note)}\n\n{body}\n")
            text = "\n---\n\n".join(parts)
        else:
            parts = []
            for note, mapping in zip(notes, mappings, strict=True):
                lines = [note.display_title, stamp(note), ""]
                for block, ordinal in numbered_blocks(note.blocks):
                    if block.kind == "table":
                        lines.extend("\t".join(row) for row in block.rows)
                    elif block.kind == "divider":
                        lines.append("─" * 20)
                    elif block.kind == "todo":
                        lines.append(("☑ " if block.checked else "☐ ") + block.text)
                    elif block.kind == "list":
                        lines.append(
                            ("  " * (block.level - 1))
                            + (f"{ordinal}. " if block.ordered else "• ")
                            + block.text
                        )
                    elif block.kind != "attachment":
                        lines.append(block.text)
                for asset in note.attachments:
                    lines.append(
                        f"[附件] {asset.name}：{mapping[asset.id][0] if asset.id in mapping else '未获取'}"
                    )
                lines.extend("[提示] " + warning for warning in note.warnings)
                parts.append("\n".join(lines))
            text = ("\n\n" + "═" * 40 + "\n\n").join(parts)
        path.write_text(text, encoding="utf-8-sig" if format == "txt" else "utf-8")

    def _docx(
        self, path: Path, notes: list[NoteDocument], mappings: list[dict], toc: bool, issues: list[ItemIssue]
    ):

        doc = Document()
        normal = doc.styles["Normal"]
        normal.font.name = "Microsoft YaHei"
        normal.font.size = Pt(10.5)
        normal.element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        if toc:
            doc.add_heading("笔记目录", 0)
            for index, note in enumerate(notes):
                self._hyperlink(doc.add_paragraph(), note.display_title, f"note_{index}", anchor=True)
            doc.add_page_break()
        for index, (note, mapping) in enumerate(zip(notes, mappings, strict=True)):
            if index:
                doc.add_page_break()
            heading = doc.add_heading(note.display_title, 1)
            start = OxmlElement("w:bookmarkStart")
            start.set(qn("w:id"), str(index))
            start.set(qn("w:name"), f"note_{index}")
            end = OxmlElement("w:bookmarkEnd")
            end.set(qn("w:id"), str(index))
            heading._p.insert(0, start)
            heading._p.append(end)
            doc.add_paragraph(stamp(note), style="Caption")
            placed = set()
            list_numbers = {}
            for block, ordinal in numbered_blocks(note.blocks):
                list_numbers = ({level: value for level, value in list_numbers.items() if level <= block.level}
                                if block.kind == "list" else {})
                if block.kind == "table":
                    if block.rows:
                        width = max(map(len, block.rows))
                        table = doc.add_table(rows=0, cols=width)
                        table.style = "Table Grid"
                        for row in block.rows:
                            cells = table.add_row().cells
                            for column, text in enumerate(row):
                                cells[column].text = text
                    continue
                if block.kind == "attachment":
                    self._docx_asset(doc, block.attachment_id, mapping, path.parent, issues, note.source_id)
                    placed.add(block.attachment_id)
                    continue
                if block.kind == "heading":
                    paragraph = doc.add_heading(level=min(9, block.level + 1))
                elif block.kind == "list":
                    paragraph = doc.add_paragraph(style="List Number" if block.ordered else "List Bullet")
                    if block.ordered:
                        if ordinal == 1 or block.level not in list_numbers:
                            numbering = doc.part.numbering_part.element
                            style_id = doc.styles["List Number"].element.pPr.numPr.numId.val
                            template = numbering.num_having_numId(style_id)
                            instance = numbering.add_num(template.abstractNumId.val)
                            instance.add_lvlOverride(0).add_startOverride(1)
                            list_numbers[block.level] = instance.numId
                        properties = paragraph._p.get_or_add_pPr().get_or_add_numPr()
                        properties.get_or_add_numId().val = list_numbers[block.level]
                        properties.get_or_add_ilvl().val = 0
                    else:
                        list_numbers.pop(block.level, None)
                else:
                    paragraph = doc.add_paragraph()
                if block.kind in ("quote", "list"):
                    paragraph.paragraph_format.left_indent = Inches(0.2 * block.level)
                if block.kind == "divider":
                    paragraph.add_run("─" * 35)
                elif block.kind == "todo":
                    paragraph.add_run("☑ " if block.checked else "☐ ")
                for span in block.spans:
                    if safe_link(span.link):
                        run = self._hyperlink(paragraph, span.text, span.link)
                    else:
                        run = paragraph.add_run(span.text)
                    run.bold, run.italic, run.underline = span.bold, span.italic, span.underline
                    run.font.strike = span.strike
                    if span.highlight:
                        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
                    if span.code or block.kind == "code":
                        run.font.name = "Consolas"
                if block.kind == "code":
                    paragraph.paragraph_format.space_after = Pt(12)
            for asset in note.attachments:
                if asset.id not in placed:
                    self._docx_asset(doc, asset.id, mapping, path.parent, issues, note.source_id)
            for warning in note.warnings:
                doc.add_paragraph("提示：" + warning)
        doc.save(path)

    @staticmethod
    def _hyperlink(paragraph, text: str, target: str, anchor: bool = False):
        from docx.opc.constants import RELATIONSHIP_TYPE

        link = OxmlElement("w:hyperlink")
        if anchor:
            link.set(qn("w:anchor"), target)
        else:
            relationship = paragraph.part.relate_to(target, RELATIONSHIP_TYPE.HYPERLINK, is_external=True)
            link.set(qn("r:id"), relationship)
        run, props = OxmlElement("w:r"), OxmlElement("w:rPr")
        color = OxmlElement("w:color")
        color.set(qn("w:val"), "6556CC")
        props.append(color)
        run.append(props)
        link.append(run)
        paragraph._p.append(link)
        formatted = Run(run, paragraph)
        formatted.text = text
        return formatted

    def _docx_asset(
        self, doc, asset_id: str | None, mapping: dict, root: Path, issues: list[ItemIssue], note_id: str
    ):
        from urllib.parse import unquote

        from docx.image.exceptions import UnrecognizedImageError

        asset = mapping.get(asset_id)
        if not asset:
            doc.add_paragraph("[附件未获取]")
            return
        relative, name, kind = asset
        if kind == "image":
            try:
                from PIL import Image

                source = confined(root, unquote(relative))
                with Image.open(source) as image:
                    width = min(6.1, image.width / 96)
                doc.add_picture(str(source), width=Inches(width))
                doc.add_paragraph(name, style="Caption")
                return
            except (OSError, ValueError, KeyError, UnrecognizedImageError):
                doc.add_paragraph("此图片格式无法内嵌，原文件已保留：")
                issues.append(
                    ItemIssue(
                        note_id=note_id,
                        code="image_not_embedded",
                        message=f"图片无法内嵌到 Word，原文件已保留：{name}",
                    )
                )
        self._hyperlink(doc.add_paragraph(), name, relative)


def package_export(source: Path, destination: Path, check_cancel: Callable[[], None] = lambda: None) -> Path:
    source, destination = source.resolve(), destination.resolve()
    if not (source / "导出说明.json").is_file():
        raise BridgeError("invalid_bundle", "请先完成一次导出，再打包本次导出文件。")
    if not destination.is_dir() or destination.is_relative_to(source):
        raise BridgeError("invalid_directory", "请选择导出目录之外的打包位置。")
    target = unused_path(destination, safe_filename(source.name, limit=140) + ".zip")
    temporary = confined(destination, f".note-bridge-{uuid.uuid4().hex}.zip")
    try:
        with zipfile.ZipFile(temporary, "w", zipfile.ZIP_DEFLATED) as archive:
            for file in sorted(source.rglob("*")):
                check_cancel()
                if file.is_symlink() or not file.resolve().is_relative_to(source):
                    raise BridgeError("unsafe_path", "导出目录包含指向外部的链接，无法打包。")
                if file.is_file():
                    archive.write(file, file.relative_to(source).as_posix())
        check_cancel()
        os.replace(temporary, target)
        return target
    finally:
        if temporary.exists() and temporary.resolve().parent == destination:
            temporary.unlink()
